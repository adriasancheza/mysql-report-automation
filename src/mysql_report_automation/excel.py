"""Build formatted Excel workbooks from query results using openpyxl."""

from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from mysql_report_automation.config import SheetConfig
from mysql_report_automation.db import QueryResult

HEADER_FILL = PatternFill(start_color="FF2F5496", end_color="FF2F5496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 8

_NUMERIC_TYPES = (int, float)


def _is_numeric_column(rows: list[tuple[Any, ...]], col_index: int) -> bool:
    values = [row[col_index] for row in rows if row[col_index] is not None]
    if not values:
        return False
    return all(isinstance(v, _NUMERIC_TYPES) and not isinstance(v, bool) for v in values)


def _is_date_column(rows: list[tuple[Any, ...]], col_index: int) -> bool:
    values = [row[col_index] for row in rows if row[col_index] is not None]
    if not values:
        return False
    return all(isinstance(v, dt.date | dt.datetime) for v in values)


def _write_header(ws: Worksheet, columns: list[str]) -> None:
    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT


def _write_rows(ws: Worksheet, rows: list[tuple[Any, ...]]) -> None:
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def _apply_column_formats(
    ws: Worksheet,
    columns: list[str],
    rows: list[tuple[Any, ...]],
    sheet_config: SheetConfig,
) -> None:
    for col_idx, name in enumerate(columns, start=1):
        fmt = sheet_config.number_formats.get(name) or sheet_config.date_formats.get(name)
        if not fmt and _is_date_column(rows, col_idx - 1):
            fmt = "yyyy-mm-dd"
        if not fmt and _is_numeric_column(rows, col_idx - 1):
            fmt = "#,##0.00"
        if fmt:
            for row_idx in range(2, len(rows) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = fmt


def _auto_size_columns(ws: Worksheet, columns: list[str], rows: list[tuple[Any, ...]]) -> None:
    for col_idx, name in enumerate(columns, start=1):
        longest = len(str(name))
        for row in rows:
            value = row[col_idx - 1]
            if value is not None:
                longest = max(longest, len(str(value)))
        width = max(MIN_COLUMN_WIDTH, min(MAX_COLUMN_WIDTH, longest + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _add_summary_sheet(
    wb: Workbook,
    columns: list[str],
    rows: list[tuple[Any, ...]],
    totals_columns: list[str],
    report_name: str,
) -> None:
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Report"
    ws["B1"] = report_name
    ws["A2"] = "Generated at"
    ws["B2"] = dt.datetime.now().isoformat(timespec="seconds")
    ws["A3"] = "Row count"
    ws["B3"] = len(rows)

    ws["A1"].font = Font(bold=True)
    ws["A2"].font = Font(bold=True)
    ws["A3"].font = Font(bold=True)

    header_row = 5
    ws.cell(row=header_row, column=1, value="Metric").font = Font(bold=True)
    ws.cell(row=header_row, column=2, value="Total").font = Font(bold=True)

    for offset, col_name in enumerate(totals_columns, start=1):
        if col_name not in columns:
            continue
        col_index = columns.index(col_name)
        total = sum(
            row[col_index]
            for row in rows
            if isinstance(row[col_index], _NUMERIC_TYPES) and not isinstance(row[col_index], bool)
        )
        ws.cell(row=header_row + offset, column=1, value=col_name)
        cell = ws.cell(row=header_row + offset, column=2, value=total)
        cell.number_format = "#,##0.00"

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 24


def build_workbook(
    result: QueryResult, sheet_config: SheetConfig, report_name: str
) -> Workbook:
    """Build a formatted openpyxl Workbook from a query result."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_config.title[:31] or "Report"

    _write_header(ws, result.columns)
    _write_rows(ws, result.rows)
    _apply_column_formats(ws, result.columns, result.rows, sheet_config)
    _auto_size_columns(ws, result.columns, result.rows)

    if result.rows:
        if sheet_config.freeze_header:
            ws.freeze_panes = "A2"
        if sheet_config.autofilter:
            last_col = get_column_letter(len(result.columns))
            ws.auto_filter.ref = f"A1:{last_col}{len(result.rows) + 1}"

    if sheet_config.summary.enabled:
        _add_summary_sheet(wb, result.columns, result.rows, sheet_config.summary.totals, report_name)

    return wb


def save_xlsx(workbook: Workbook, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def save_csv(result: QueryResult, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(result.columns)
        writer.writerows(result.rows)
    return path
