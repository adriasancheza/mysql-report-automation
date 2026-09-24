from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from mysql_report_automation.config import SheetConfig, SummaryConfig
from mysql_report_automation.db import QueryResult
from mysql_report_automation.excel import build_workbook, save_csv, save_xlsx


def _sample_result() -> QueryResult:
    return QueryResult(
        columns=["region", "order_date", "units", "revenue"],
        rows=[
            ("North", date(2026, 1, 5), 10, 100.5),
            ("South", date(2026, 1, 6), 3, 90.25),
        ],
    )


def test_build_workbook_header_and_rows():
    result = _sample_result()
    sheet = SheetConfig(title="Sales")
    wb = build_workbook(result, sheet, "test_report")
    ws = wb.active

    assert ws.title == "Sales"
    assert [c.value for c in ws[1]] == result.columns
    assert ws.cell(row=2, column=1).value == "North"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:D3"


def test_build_workbook_number_and_date_formats_applied():
    result = _sample_result()
    sheet = SheetConfig(
        title="Sales",
        number_formats={"revenue": "#,##0.00"},
        date_formats={"order_date": "yyyy-mm-dd"},
    )
    wb = build_workbook(result, sheet, "test_report")
    ws = wb.active

    assert ws.cell(row=2, column=4).number_format == "#,##0.00"
    assert ws.cell(row=2, column=2).number_format == "yyyy-mm-dd"


def test_summary_sheet_created_with_totals():
    result = _sample_result()
    sheet = SheetConfig(
        title="Sales", summary=SummaryConfig(enabled=True, totals=["units", "revenue"])
    )
    wb = build_workbook(result, sheet, "test_report")

    assert "Summary" in wb.sheetnames
    summary = wb["Summary"]
    values = {summary.cell(row=r, column=1).value: summary.cell(row=r, column=2).value for r in range(6, 8)}
    assert values["units"] == 13
    assert values["revenue"] == 190.75


def test_empty_result_does_not_crash():
    result = QueryResult(columns=["a", "b"], rows=[])
    wb = build_workbook(result, SheetConfig(), "empty")
    ws = wb.active
    assert [c.value for c in ws[1]] == ["a", "b"]
    assert ws.freeze_panes is None


def test_save_xlsx_and_csv_roundtrip(tmp_path):
    result = _sample_result()
    wb = build_workbook(result, SheetConfig(title="Sales"), "test_report")

    xlsx_path = save_xlsx(wb, tmp_path / "out" / "report.xlsx")
    assert xlsx_path.exists()
    reloaded = load_workbook(xlsx_path)
    assert reloaded.active.cell(row=2, column=1).value == "North"

    csv_path = save_csv(result, tmp_path / "out" / "report.csv")
    text = csv_path.read_text(encoding="utf-8")
    assert "region,order_date,units,revenue" in text
    assert "North" in text
